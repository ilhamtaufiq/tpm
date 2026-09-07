import openpyxl
wb = openpyxl.load_workbook(
    'C:/Users/asusg/Downloads/TPM_IMPORT_TEMPLATE_REAL PERIODE BERJALAN.xlsx',
    data_only=True,
)
for name in ['customers', 'suppliers', 'jasa_servis', 'karyawan',
             'kas_opening', 'asset', 'hutang_opening', 'piutang_opening',
             'armada', 'supir', 'mobil']:
    ws = wb[name]
    print(f'=== {name} ({ws.max_row - 1} data rows) ===')
    for row in ws.iter_rows(min_row=2, values_only=True):
        print('  ', row)
